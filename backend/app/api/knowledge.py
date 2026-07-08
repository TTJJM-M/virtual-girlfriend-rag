import os
import asyncio
import aiofiles
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, delete
from app.database import get_async_db, AsyncSessionLocal
from app.models.knowledge import Document, DocumentChunk, DocumentStatus
from app.models.user import User
from app.core.deps import get_current_active_user
from app.schemas.knowledge import DocumentResponse, DocumentListResponse, DocumentProcessStatus
from app.services.vector_store import get_vector_store
from app.config import settings
from typing import Annotated, List
from langchain_text_splitters import RecursiveCharacterTextSplitter

router = APIRouter(prefix="/api/knowledge", tags=["知识库管理"])


def get_file_type(filename: str) -> str:
    """Get file type from extension"""
    ext = filename.split(".")[-1].lower()
    if ext in ["txt", "md", "markdown", "pdf", "doc", "docx"]:
        return ext
    return "txt"


def detect_file_encoding(file_path: str) -> str:
    """Auto-detect file encoding, trying UTF-8 first, then Chinese encodings"""
    encodings = ["utf-8", "gbk", "gb2312", "gb18030", "latin-1"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                f.read()
                return enc
        except (UnicodeDecodeError, LookupError):
            continue
    return "utf-8"


async def read_file_with_encoding(file_path: str, encoding: str = None) -> str:
    """Read file content with auto-detected encoding"""
    if encoding is None:
        encoding = detect_file_encoding(file_path)
    async with aiofiles.open(file_path, "r", encoding=encoding) as f:
        return await f.read()


async def parse_document(file_path: str, file_type: str) -> str:
    """Parse document content based on file type"""
    content = ""

    if file_type in ["txt", "md", "markdown"]:
        encoding = detect_file_encoding(file_path)
        async with aiofiles.open(file_path, "r", encoding=encoding) as f:
            content = await f.read()

    elif file_type == "pdf":
        # PDF is binary — must use pdfplumber, NEVER read as text
        import pdfplumber
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    content += text + "\n"

    elif file_type in ["doc", "docx"]:
        # DOCX is binary — must use python-docx, NEVER read as text
        from docx import Document as DocxDocument
        doc = DocxDocument(file_path)
        for para in doc.paragraphs:
            content += para.text + "\n"

    return content.strip()


@router.post("/upload", response_model=DocumentResponse)
async def upload_document(
    file: UploadFile = File(...),
    current_user: Annotated[User, Depends(get_current_active_user)] = None,
    db: Annotated[AsyncSession, Depends(get_async_db)] = None,
):
    """Upload and process a document"""
    content = await file.read()
    if len(content) > settings.MAX_UPLOAD_SIZE:
        raise HTTPException(status_code=400, detail=f"文件大小超过限制 ({settings.MAX_UPLOAD_SIZE // 1024 // 1024}MB)")

    file_type = get_file_type(file.filename)
    file_path = os.path.join(settings.UPLOAD_DIR, str(current_user.id), f"{file.filename}")
    os.makedirs(os.path.dirname(file_path), exist_ok=True)

    async with aiofiles.open(file_path, "wb") as f:
        await f.write(content)

    doc = Document(
        user_id=current_user.id,
        filename=file.filename,
        file_path=file_path,
        file_size=len(content),
        file_type=file_type,
        status=DocumentStatus.PARSING,
    )
    db.add(doc)
    await db.commit()
    await db.refresh(doc)

    # Process document in background with its own DB session
    asyncio.create_task(process_document(doc.id, file_path, file_type, current_user.id))

    return DocumentResponse(
        id=doc.id,
        filename=doc.filename,
        file_size=doc.file_size,
        file_type=doc.file_type,
        status=doc.status,
        chunk_count=doc.chunk_count,
        error_message=doc.error_message,
        created_at=doc.created_at,
        updated_at=doc.updated_at,
    )


async def process_document(
    doc_id: int,
    file_path: str,
    file_type: str,
    user_id: int,
):
    """Process document in background with its own DB session"""
    async with AsyncSessionLocal() as db:
        try:
            result = await db.execute(
                select(Document).where(Document.id == doc_id)
            )
            doc = result.scalar_one_or_none()
            if not doc:
                return

            # Parse
            doc.status = DocumentStatus.PARSING
            await db.commit()
            content = await parse_document(file_path, file_type)

            if not content:
                doc.status = DocumentStatus.FAILED
                doc.error_message = "文档内容为空或无法解析"
                await db.commit()
                return

            # Chunk
            doc.status = DocumentStatus.CHUNKING
            await db.commit()

            splitter = RecursiveCharacterTextSplitter(
                chunk_size=settings.CHUNK_SIZE,
                chunk_overlap=settings.CHUNK_OVERLAP,
                separators=["\n\n", "\n", "。", "！", "？", ".", "!", "?", " "],
            )
            chunks = splitter.split_text(content)

            # Save chunks to database
            saved_chunks = []
            for i, chunk_text in enumerate(chunks):
                chunk = DocumentChunk(
                    document_id=doc_id,
                    chunk_index=i,
                    content=chunk_text,
                    token_count=len(chunk_text),
                )
                db.add(chunk)
                saved_chunks.append(chunk)

            await db.commit()

            # Embed and index
            doc.status = DocumentStatus.EMBEDDING
            doc.chunk_count = len(chunks)
            await db.commit()

            vector_store = get_vector_store(user_id)
            chunk_data = [
                (chunk.content, chunk.chunk_index, chunk.document_id, doc.filename)
                for chunk in saved_chunks
            ]
            vector_store.add_chunks(chunk_data)

            # Done
            doc.status = DocumentStatus.READY
            await db.commit()

            print(f"✓ Document {doc.filename} processed: {len(chunks)} chunks")

        except Exception as e:
            print(f"✗ Document processing failed: {e}")
            result = await db.execute(
                select(Document).where(Document.id == doc_id)
            )
            doc = result.scalar_one_or_none()
            if doc:
                doc.status = DocumentStatus.FAILED
                doc.error_message = str(e)
                await db.commit()


@router.get("", response_model=DocumentListResponse)
async def list_documents(
    current_user: Annotated[User, Depends(get_current_active_user)],
    db: Annotated[AsyncSession, Depends(get_async_db)],
):
    """List user's documents"""
    result = await db.execute(
        select(Document)
        .where(Document.user_id == current_user.id)
        .order_by(Document.created_at.desc())
    )
    documents = result.scalars().all()

    return DocumentListResponse(
        documents=[
            DocumentResponse(
                id=d.id,
                filename=d.filename,
                file_size=d.file_size,
                file_type=d.file_type,
                status=d.status,
                chunk_count=d.chunk_count,
                error_message=d.error_message,
                created_at=d.created_at,
                updated_at=d.updated_at,
            )
            for d in documents
        ],
        total=len(documents),
    )


@router.delete("/{document_id}")
async def delete_document(
    document_id: int,
    current_user: Annotated[User, Depends(get_current_active_user)],
    db: Annotated[AsyncSession, Depends(get_async_db)],
):
    """Delete a document, its chunks, and its vectors"""
    result = await db.execute(
        select(Document).where(
            Document.id == document_id,
            Document.user_id == current_user.id,
        )
    )
    doc = result.scalar_one_or_none()

    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")

    # Delete chunks from database
    await db.execute(
        delete(DocumentChunk).where(DocumentChunk.document_id == document_id)
    )

    # Delete document
    await db.delete(doc)
    await db.commit()

    # Remove from vector store
    vector_store = get_vector_store(current_user.id)
    vector_store.remove_document(document_id)

    # Delete file from disk
    try:
        os.remove(doc.file_path)
    except OSError:
        pass

    return {"message": "文档已删除"}


@router.get("/{document_id}/status", response_model=DocumentProcessStatus)
async def get_document_status(
    document_id: int,
    current_user: Annotated[User, Depends(get_current_active_user)],
    db: Annotated[AsyncSession, Depends(get_async_db)],
):
    """Get document processing status"""
    result = await db.execute(
        select(Document).where(
            Document.id == document_id,
            Document.user_id == current_user.id,
        )
    )
    doc = result.scalar_one_or_none()

    if not doc:
        raise HTTPException(status_code=404, detail="文档不存在")

    progress_map = {
        DocumentStatus.UPLOADING: 10,
        DocumentStatus.PARSING: 30,
        DocumentStatus.CHUNKING: 50,
        DocumentStatus.EMBEDDING: 80,
        DocumentStatus.READY: 100,
        DocumentStatus.FAILED: 0,
    }

    return DocumentProcessStatus(
        status=doc.status,
        chunk_count=doc.chunk_count,
        error_message=doc.error_message,
        progress=progress_map.get(doc.status, 0),
    )


@router.get("/stats")
async def get_knowledge_stats(
    current_user: Annotated[User, Depends(get_current_active_user)],
):
    """Get knowledge base statistics"""
    vector_store = get_vector_store(current_user.id)
    stats = vector_store.get_stats()

    return {
        "total_chunks": stats["total_chunks"],
        "total_vectors": stats["total_vectors"],
        "model_status": stats["model_status"],
    }
